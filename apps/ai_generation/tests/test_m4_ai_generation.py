import io
from decimal import Decimal
from unittest.mock import patch
import docx
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APITestCase
from apps.accounts.models import (
    User,
    Institute,
    UserInstituteMembership,
    RoleChoices,
    MembershipStatus,
)
from apps.academics.models import (
    AcademicYear,
    Department,
    Subject,
    TeacherAssignment,
    Unit,
    Topic,
    LearningOutcome,
)
from apps.questions.models import (
    Question,
    QuestionOrigin,
    QuestionStatus,
    QuestionType,
)
from apps.ai_generation.models import (
    SourceDocument,
    SourceDocumentChunk,
    QuestionGenerationJob,
    GeneratedCandidateRecord,
    GenerationJobStatus,
    CandidateStatus,
    DocumentProcessingStatus,
)
from apps.ai_generation.services.gemini_service import (
    GoogleGeminiProvider,
    AIConfigurationError,
    AIProviderError,
    AISchemaError,
)
from apps.ai_generation.services.orchestrator import execute_question_generation_job


def build_sample_pdf_bytes(text="Data Structures and Algorithms in Tree Traversal"):
    """Creates a minimal valid PDF byte string containing extractable text."""
    return b"""%PDF-1.4
1 0 obj <</Type /Catalog /Pages 2 0 R>> endobj
2 0 obj <</Type /Pages /Kids [3 0 R] /Count 1>> endobj
3 0 obj <</Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources <</Font <</F1 5 0 R>>>>>> endobj
4 0 obj <</Length 44>> stream
BT
/F1 12 Tf
72 712 Td
(""" + text.encode('ascii', errors='ignore') + b""") Tj
ET
endstream endobj
5 0 obj <</Type /Font /Subtype /Type1 /BaseFont /Helvetica>> endobj
xref
0 6
0000000000 65535 f 
0000000010 00000 n 
0000000060 00000 n 
0000000117 00000 n 
0000000236 00000 n 
0000000330 00000 n 
trailer <</Size 6 /Root 1 0 R>>
startxref
407
%%EOF"""


def build_sample_docx_bytes(title="AVL Tree Properties", body="An AVL tree is a self-balancing binary search tree."):
    """Creates a valid DOCX document in memory."""
    doc = docx.Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(body)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class M4AIGenerationEngineTests(APITestCase):

    def setUp(self):
        # 1. Platform Admin
        self.platform_admin = User.objects.create_superuser(
            email='super@examcraft.org',
            password='AdminPassword123!'
        )

        # 2. Institutes: Alpha & Beta
        self.institute_alpha = Institute.objects.create(code='ALPHA', name='Alpha University')
        self.institute_beta = Institute.objects.create(code='BETA', name='Beta University')
        self.academic_year_alpha = AcademicYear.objects.create(
            institute=self.institute_alpha,
            label='2026-2027',
            start_date='2026-07-01',
            end_date='2027-06-30',
        )

        # 3. Alpha Teacher
        self.teacher_alpha = User.objects.create_user(
            email='teacher@alpha.edu',
            password='Password123!',
            first_name='Barbara',
            last_name='Liskov',
            is_verified=True
        )
        UserInstituteMembership.objects.create(
            user=self.teacher_alpha,
            institute=self.institute_alpha,
            role=RoleChoices.TEACHER,
            status=MembershipStatus.ACTIVE
        )

        # 4. Alpha Student
        self.student_alpha = User.objects.create_user(
            email='student@alpha.edu',
            password='Password123!',
            first_name='Claude',
            last_name='Shannon',
            is_verified=True
        )
        UserInstituteMembership.objects.create(
            user=self.student_alpha,
            institute=self.institute_alpha,
            role=RoleChoices.STUDENT,
            status=MembershipStatus.ACTIVE
        )

        # 5. Academic Hierarchy for Alpha
        self.dept_cs = Department.objects.create(institute=self.institute_alpha, code='CSE', name='Computer Science')
        self.subj_dsa = Subject.objects.create(institute=self.institute_alpha, department=self.dept_cs, code='CS201', name='Data Structures')
        self.unit_trees = Unit.objects.create(subject=self.subj_dsa, unit_number=1, title='Balanced Trees')
        self.topic_avl = Topic.objects.create(unit=self.unit_trees, title='AVL Trees', order=1)
        self.lo_trees = LearningOutcome.objects.create(
            subject=self.subj_dsa,
            code='CO1',
            description='Analyze time complexity of self-balancing search trees.'
        )
        self.teacher_alpha_membership = self.teacher_alpha.memberships.get(
            institute=self.institute_alpha,
            role=RoleChoices.TEACHER,
        )
        TeacherAssignment.objects.create(
            teacher=self.teacher_alpha_membership,
            subject=self.subj_dsa,
            academic_year=self.academic_year_alpha,
        )

        # 6. Academic Hierarchy for Beta
        self.dept_beta = Department.objects.create(institute=self.institute_beta, code='B-MATH', name='Beta Math')
        self.subj_beta = Subject.objects.create(institute=self.institute_beta, department=self.dept_beta, code='MATH101', name='Calculus')

    # --------------------------------------------------------------------------
    # 1. Configuration & Security Tests
    # --------------------------------------------------------------------------

    def test_missing_gemini_api_key_configuration_fails_cleanly(self):
        """
        When GEMINI_API_KEY is not configured or empty, the provider raises
        AIConfigurationError and does not simulate or invent questions.
        """
        with patch('django.conf.settings.GEMINI_API_KEY', ''):
            with self.assertRaises(AIConfigurationError) as ctx:
                GoogleGeminiProvider(api_key='')
            self.assertIn("API key is missing", str(ctx.exception))

        # Test within orchestrator execution
        doc = SourceDocument.objects.create(
            institute=self.institute_alpha,
            subject=self.subj_dsa,
            title='DSA Syllabus',
            file_type='txt',
            processing_status='completed',
            extracted_text='Binary search trees properties'
        )
        SourceDocumentChunk.objects.create(
            document=doc,
            chunk_index=1,
            content='Binary search trees properties and AVL balancing'
        )

        job = QuestionGenerationJob.objects.create(
            institute=self.institute_alpha,
            subject=self.subj_dsa,
            requested_by=self.teacher_alpha,
            num_questions_requested=3,
            question_types=['mcq']
        )
        job.source_documents.add(doc)

        with patch('django.conf.settings.GEMINI_API_KEY', ''):
            executed_job = execute_question_generation_job(job)
            self.assertEqual(executed_job.status, GenerationJobStatus.FAILED)
            self.assertIn("API key is missing", executed_job.error_message)
            self.assertEqual(executed_job.num_questions_generated, 0)
            self.assertEqual(Question.objects.count(), 0)

    # --------------------------------------------------------------------------
    # 2. Document Upload & Text Extraction Tests
    # --------------------------------------------------------------------------

    def test_unsupported_file_extension_upload_rejected(self):
        self.client.force_authenticate(user=self.teacher_alpha)
        url = reverse('source-material-list')
        fake_exe = SimpleUploadedFile("script.exe", b"binary content", content_type="application/octet-stream")
        data = {
            'institute': str(self.institute_alpha.id),
            'subject': str(self.subj_dsa.id),
            'title': 'Malicious file attempt',
            'file': fake_exe
        }
        response = self.client.post(url, data, format='multipart')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('file', response.data['error']['details'])

    def test_empty_file_upload_rejected(self):
        self.client.force_authenticate(user=self.teacher_alpha)
        url = reverse('source-material-list')
        empty_txt = SimpleUploadedFile("empty.txt", b"", content_type="text/plain")
        data = {
            'institute': str(self.institute_alpha.id),
            'subject': str(self.subj_dsa.id),
            'title': 'Empty notes',
            'file': empty_txt
        }
        response = self.client.post(url, data, format='multipart')
        # Uploading 0 byte file is marked failed or rejected
        if response.status_code == status.HTTP_201_CREATED:
            self.assertEqual(response.data['processing_status'], 'failed')
            self.assertIn("empty", response.data['error_message'].lower())
        else:
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_text_extraction_from_pdf(self):
        self.client.force_authenticate(user=self.teacher_alpha)
        url = reverse('source-material-list')
        pdf_bytes = build_sample_pdf_bytes("Balanced Trees and Rotations")
        uploaded_pdf = SimpleUploadedFile("trees_chapter.pdf", pdf_bytes, content_type="application/pdf")

        data = {
            'institute': str(self.institute_alpha.id),
            'subject': str(self.subj_dsa.id),
            'unit': str(self.unit_trees.id),
            'topic': str(self.topic_avl.id),
            'title': 'Chapter 1: Balanced Trees PDF',
            'file': uploaded_pdf
        }
        response = self.client.post(url, data, format='multipart')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['processing_status'], 'completed')
        self.assertEqual(response.data['file_type'], 'pdf')
        self.assertGreater(response.data['chunks_count'], 0)

        # Check chunks endpoint
        doc_id = response.data['id']
        chunks_url = reverse('source-material-chunks', kwargs={'pk': doc_id})
        chunks_resp = self.client.get(chunks_url)
        self.assertEqual(chunks_resp.status_code, status.HTTP_200_OK)
        self.assertGreater(len(chunks_resp.data['results']), 0)
        self.assertIn("Balanced Trees", chunks_resp.data['results'][0]['content'])

    def test_text_extraction_from_docx(self):
        self.client.force_authenticate(user=self.teacher_alpha)
        url = reverse('source-material-list')
        docx_bytes = build_sample_docx_bytes("Red-Black Trees Invariant", "Each node is either red or black.")
        uploaded_docx = SimpleUploadedFile("rbtree.docx", docx_bytes, content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

        data = {
            'institute': str(self.institute_alpha.id),
            'subject': str(self.subj_dsa.id),
            'title': 'Red-Black Tree Notes',
            'file': uploaded_docx
        }
        response = self.client.post(url, data, format='multipart')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['processing_status'], 'completed')
        self.assertEqual(response.data['file_type'], 'docx')
        self.assertIn("Red-Black Trees", response.data['chunks_count'] > 0 and SourceDocument.objects.get(id=response.data['id']).extracted_text)

    def test_text_extraction_from_txt(self):
        self.client.force_authenticate(user=self.teacher_alpha)
        url = reverse('source-material-list')
        txt_content = b"Graph algorithms: BFS uses a queue, DFS uses a recursion stack."
        uploaded_txt = SimpleUploadedFile("graphs.txt", txt_content, content_type="text/plain")

        data = {
            'institute': str(self.institute_alpha.id),
            'subject': str(self.subj_dsa.id),
            'title': 'Graph Traversal Syllabus',
            'file': uploaded_txt
        }
        response = self.client.post(url, data, format='multipart')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['processing_status'], 'completed')
        self.assertEqual(response.data['file_type'], 'txt')

    # --------------------------------------------------------------------------
    # 3. Permissions & Tenant Isolation
    # --------------------------------------------------------------------------

    def test_student_role_forbidden_from_uploads_and_generation(self):
        self.client.force_authenticate(user=self.student_alpha)

        # Upload attempt
        upload_url = reverse('source-material-list')
        txt_file = SimpleUploadedFile("notes.txt", b"Notes", content_type="text/plain")
        up_resp = self.client.post(upload_url, {
            'institute': str(self.institute_alpha.id),
            'subject': str(self.subj_dsa.id),
            'title': 'Student notes',
            'file': txt_file
        }, format='multipart')
        self.assertEqual(up_resp.status_code, status.HTTP_403_FORBIDDEN)

        # Generation job attempt
        job_url = reverse('question-generation-job-list')
        gen_resp = self.client.post(job_url, {
            'institute': str(self.institute_alpha.id),
            'subject': str(self.subj_dsa.id),
            'num_questions': 5
        }, format='json')
        self.assertEqual(gen_resp.status_code, status.HTTP_403_FORBIDDEN)

    def test_source_material_tenant_isolation(self):
        doc_beta = SourceDocument.objects.create(
            institute=self.institute_beta,
            subject=self.subj_beta,
            title='Beta Confidential Calculus Material',
            file_type='txt',
            processing_status='completed',
            extracted_text='Confidential beta calculus notes'
        )

        self.client.force_authenticate(user=self.teacher_alpha)
        detail_url = reverse('source-material-detail', kwargs={'pk': doc_beta.id})
        response = self.client.get(detail_url)
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

        # Attempt to use Beta document in Alpha generation job
        job_url = reverse('question-generation-job-list')
        job_payload = {
            'institute': str(self.institute_alpha.id),
            'subject': str(self.subj_dsa.id),
            'source_document_ids': [str(doc_beta.id)],
            'num_questions': 3
        }
        job_resp = self.client.post(job_url, job_payload, format='json')
        self.assertEqual(job_resp.status_code, status.HTTP_400_BAD_REQUEST)

        unauthorized_job = self.client.post(job_url, {
            'institute': str(self.institute_beta.id),
            'subject': str(self.subj_beta.id),
            'num_questions': 2,
        }, format='json')
        self.assertEqual(unauthorized_job.status_code, status.HTTP_403_FORBIDDEN)

    def test_request_question_count_limits(self):
        self.client.force_authenticate(user=self.teacher_alpha)
        job_url = reverse('question-generation-job-list')
        payload = {
            'institute': str(self.institute_alpha.id),
            'subject': str(self.subj_dsa.id),
            'num_questions': 50  # Exceeds max 20 limit
        }
        response = self.client.post(job_url, payload, format='json')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    # --------------------------------------------------------------------------
    # 4. Mocked Grounded Generation & Validation Tests
    # --------------------------------------------------------------------------

    @patch('apps.ai_generation.services.orchestrator.GoogleGeminiProvider')
    def test_successful_grounded_generation_with_mocked_provider(self, MockProvider):
        """
        Validates creation of M3 draft questions (MCQ, True/False, Short, Long, Coding)
        with normalized options, rubrics, coding test cases, and provenance.
        """
        # Create processed document in Alpha
        doc = SourceDocument.objects.create(
            institute=self.institute_alpha,
            subject=self.subj_dsa,
            unit=self.unit_trees,
            topic=self.topic_avl,
            title='AVL Trees Reference Manual',
            file_type='txt',
            processing_status='completed',
            extracted_text='An AVL tree maintains height balance where left and right subtrees differ in height by at most 1.'
        )
        SourceDocumentChunk.objects.create(
            document=doc,
            chunk_index=1,
            page_number=1,
            content='An AVL tree maintains height balance where left and right subtrees differ in height by at most 1. Rebalancing is done via LL, RR, LR, RL rotations.'
        )

        mocked_candidates = [
            # 1. MCQ Candidate
            {
                'question_type': 'mcq',
                'question_text': 'What is the maximum allowed height difference between subtrees of any node in an AVL tree?',
                'difficulty': 'easy',
                'bloom_taxonomy': 'remember',
                'marks': 1.0,
                'negative_marks': 0.25,
                'options': [
                    {'option_key': 'A', 'option_text': '0', 'is_correct': False},
                    {'option_key': 'B', 'option_text': '1', 'is_correct': True, 'explanation': 'By AVL definition |h_left - h_right| <= 1'},
                    {'option_key': 'C', 'option_text': '2', 'is_correct': False},
                    {'option_key': 'D', 'option_text': 'No limit', 'is_correct': False},
                ],
                'provenance': {
                    'source_title': 'AVL Trees Reference Manual',
                    'page_or_section': 'Page 1',
                    'excerpt': 'An AVL tree maintains height balance where left and right subtrees differ in height by at most 1.',
                    'relevance_note': 'Defines the AVL balance factor condition'
                },
                'supported_by_source': True
            },
            # 2. True / False Candidate
            {
                'question_type': 'true_false',
                'question_text': 'Rebalancing in an AVL tree involves rotations such as LL, RR, LR, and RL.',
                'difficulty': 'easy',
                'bloom_taxonomy': 'understand',
                'marks': 1.0,
                'options': [
                    {'option_key': 'A', 'option_text': 'True', 'is_correct': True},
                    {'option_key': 'B', 'option_text': 'False', 'is_correct': False},
                ],
                'provenance': {
                    'source_title': 'AVL Trees Reference Manual',
                    'page_or_section': 'Page 1',
                    'excerpt': 'Rebalancing is done via LL, RR, LR, RL rotations.',
                    'relevance_note': 'States rotation types'
                },
                'supported_by_source': True
            },
            # 3. Short Answer Candidate
            {
                'question_type': 'short_answer',
                'question_text': 'List the four rotation operations used to restore height balance in an AVL tree.',
                'difficulty': 'medium',
                'bloom_taxonomy': 'remember',
                'marks': 2.0,
                'answer_key': {
                    'expected_answer': 'LL rotation, RR rotation, LR rotation, and RL rotation.',
                    'sample_solution': 'Single rotations (LL, RR) and double rotations (LR, RL).',
                    'keywords': ['LL', 'RR', 'LR', 'RL', 'rotation']
                },
                'provenance': {
                    'source_title': 'AVL Trees Reference Manual',
                    'page_or_section': 'Page 1',
                    'excerpt': 'Rebalancing is done via LL, RR, LR, RL rotations.',
                    'relevance_note': 'Direct citation of rotation operations'
                },
                'supported_by_source': True
            },
            # 4. Long Answer with Rubrics Candidate
            {
                'question_type': 'long_answer',
                'question_text': 'Explain the AVL tree balance factor and discuss when an LL rotation is required.',
                'difficulty': 'hard',
                'bloom_taxonomy': 'analyze',
                'marks': 5.0,
                'answer_key': {
                    'expected_answer': 'Balance factor is height(left) - height(right). LL rotation is required when balance factor becomes +2 due to left-left insertion.',
                    'sample_solution': 'Detailed explanation of balance factor and right rotation.',
                    'keywords': ['balance factor', 'LL rotation', 'right rotation']
                },
                'rubrics': [
                    {'criterion': 'Balance factor definition', 'description': 'Formula height(L) - height(R)', 'max_marks': 2.0, 'order': 1},
                    {'criterion': 'LL rotation trigger', 'description': 'Explains left-heavy condition', 'max_marks': 3.0, 'order': 2},
                ],
                'provenance': {
                    'source_title': 'AVL Trees Reference Manual',
                    'page_or_section': 'Page 1',
                    'excerpt': 'An AVL tree maintains height balance where left and right subtrees differ in height by at most 1.',
                    'relevance_note': 'Detailed analysis of balancing conditions'
                },
                'supported_by_source': True
            },
            # 5. Coding Question Candidate
            {
                'question_type': 'coding',
                'question_text': 'Implement a function `get_balance_factor(node)` that calculates the balance factor of a binary tree node.',
                'difficulty': 'medium',
                'bloom_taxonomy': 'apply',
                'marks': 10.0,
                'code_metadata': {
                    'language': 'python',
                    'starter_template': 'def get_balance_factor(node):\n    pass'
                },
                'answer_key': {
                    'expected_answer': 'def get_balance_factor(node): return height(node.left) - height(node.right)',
                    'test_cases': [
                        {'input': 'Node(left_height=2, right_height=1)', 'expected_output': '1', 'is_hidden': False},
                        {'input': 'Node(left_height=1, right_height=3)', 'expected_output': '-2', 'is_hidden': True},
                    ]
                },
                'provenance': {
                    'source_title': 'AVL Trees Reference Manual',
                    'page_or_section': 'Page 1',
                    'excerpt': 'An AVL tree maintains height balance where left and right subtrees differ in height by at most 1.',
                    'relevance_note': 'Coding exercise on balance factor calculation'
                },
                'supported_by_source': True
            }
        ]

        mock_instance = MockProvider.return_value
        mock_instance.generate_questions.return_value = (
            mocked_candidates,
            {'provider': 'mock_gemini', 'model': 'gemini-3.8-flash', 'candidates_returned': 5}
        )

        self.client.force_authenticate(user=self.teacher_alpha)
        job_url = reverse('question-generation-job-list')
        payload = {
            'institute': str(self.institute_alpha.id),
            'subject': str(self.subj_dsa.id),
            'unit': str(self.unit_trees.id),
            'topic': str(self.topic_avl.id),
            'source_document_ids': [str(doc.id)],
            'question_types': ['mcq', 'true_false', 'short_answer', 'long_answer', 'coding'],
            'num_questions': 5
        }

        response = self.client.post(job_url, payload, format='json')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['status'], 'completed')
        self.assertEqual(response.data['num_questions_generated'], 5)
        self.assertEqual(response.data['num_questions_failed_validation'], 0)

        # Verify created M3 Questions in database
        questions = Question.objects.filter(subject=self.subj_dsa)
        self.assertEqual(questions.count(), 5)

        for q in questions:
            # Must strictly be AI_GENERATED and DRAFT
            self.assertEqual(q.origin_type, QuestionOrigin.AI_GENERATED)
            self.assertEqual(q.status, QuestionStatus.DRAFT)
            self.assertEqual(q.current_version, 1)
            self.assertEqual(q.created_by, self.teacher_alpha)

            # Must have source reference provenance attached
            self.assertTrue(q.source_references.exists())
            ref = q.source_references.first()
            self.assertEqual(ref.source_title, 'AVL Trees Reference Manual')
            self.assertTrue(len(ref.excerpt) > 10)

        # Check candidate diagnostics endpoint
        job_id = response.data['id']
        cand_url = reverse('question-generation-job-candidates', kwargs={'pk': job_id})
        cand_resp = self.client.get(cand_url)
        self.assertEqual(cand_resp.status_code, status.HTTP_200_OK)
        self.assertEqual(len(cand_resp.data), 5)
        for c in cand_resp.data:
            self.assertEqual(c['status'], 'saved')
            self.assertIsNotNone(c['persisted_question_id'])

    @patch('apps.ai_generation.services.orchestrator.GoogleGeminiProvider')
    def test_malformed_json_and_schema_invalid_provider_response(self, MockProvider):
        doc = SourceDocument.objects.create(
            institute=self.institute_alpha,
            subject=self.subj_dsa,
            title='DSA Syllabus',
            file_type='txt',
            processing_status='completed',
            extracted_text='Binary search trees properties'
        )
        SourceDocumentChunk.objects.create(document=doc, chunk_index=1, content='Content')

        mock_instance = MockProvider.return_value
        mock_instance.generate_questions.side_effect = AISchemaError("Model response is not valid JSON")

        self.client.force_authenticate(user=self.teacher_alpha)
        job_url = reverse('question-generation-job-list')
        payload = {
            'institute': str(self.institute_alpha.id),
            'subject': str(self.subj_dsa.id),
            'source_document_ids': [str(doc.id)],
            'num_questions': 3
        }

        response = self.client.post(job_url, payload, format='json')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['status'], 'failed')
        self.assertIn("not valid JSON", response.data['error_message'])
        self.assertEqual(Question.objects.count(), 0)

    @patch('apps.ai_generation.services.orchestrator.GoogleGeminiProvider')
    def test_candidate_with_unsupported_content_or_missing_evidence(self, MockProvider):
        """
        Candidates with supported_by_source=False or missing provenance excerpt
        must NOT be persisted into M3 Question Bank.
        """
        doc = SourceDocument.objects.create(
            institute=self.institute_alpha,
            subject=self.subj_dsa,
            title='Tree Notes',
            file_type='txt',
            processing_status='completed',
            extracted_text='Binary search trees properties'
        )
        SourceDocumentChunk.objects.create(document=doc, chunk_index=1, content='Content')

        mocked_candidates = [
            # Candidate 1: Supported by source is False
            {
                'question_type': 'mcq',
                'question_text': 'Unsupported quantum computing question in a DSA exam?',
                'marks': 1.0,
                'options': [
                    {'option_key': 'A', 'option_text': '1', 'is_correct': True},
                    {'option_key': 'B', 'option_text': '2', 'is_correct': False},
                ],
                'supported_by_source': False,  # UNSUPPORTED
                'provenance': {'excerpt': '', 'source_title': ''}
            },
            # Candidate 2: Missing provenance excerpt
            {
                'question_type': 'mcq',
                'question_text': 'Hallucinated question with no citation passage?',
                'marks': 1.0,
                'options': [
                    {'option_key': 'A', 'option_text': '1', 'is_correct': True},
                    {'option_key': 'B', 'option_text': '2', 'is_correct': False},
                ],
                'supported_by_source': True,
                'provenance': {'excerpt': '', 'source_title': 'Tree Notes'}  # EMPTY EXCERPT
            }
        ]

        mock_instance = MockProvider.return_value
        mock_instance.generate_questions.return_value = (mocked_candidates, {})

        self.client.force_authenticate(user=self.teacher_alpha)
        job_url = reverse('question-generation-job-list')
        payload = {
            'institute': str(self.institute_alpha.id),
            'subject': str(self.subj_dsa.id),
            'source_document_ids': [str(doc.id)],
            'num_questions': 2
        }

        response = self.client.post(job_url, payload, format='json')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['status'], 'failed')
        self.assertEqual(response.data['num_questions_generated'], 0)
        self.assertEqual(response.data['num_questions_failed_validation'], 2)

        # Zero questions persisted in M3
        self.assertEqual(Question.objects.count(), 0)

        # Diagnostics show why they failed
        job_id = response.data['id']
        candidates = GeneratedCandidateRecord.objects.filter(job_id=job_id)
        self.assertEqual(candidates.count(), 2)
        self.assertEqual(candidates[0].status, CandidateStatus.UNSUPPORTED_CONTENT)
        self.assertEqual(candidates[1].status, CandidateStatus.VALIDATION_FAILED)

    @patch('apps.ai_generation.services.orchestrator.GoogleGeminiProvider')
    def test_invalid_candidate_marks_and_options_rejected(self, MockProvider):
        """
        Invalid MCQ (no correct option) and negative marks > positive marks
        must be caught by validator and rejected.
        """
        doc = SourceDocument.objects.create(
            institute=self.institute_alpha,
            subject=self.subj_dsa,
            title='Notes',
            file_type='txt',
            processing_status='completed',
            extracted_text='Text'
        )
        SourceDocumentChunk.objects.create(document=doc, chunk_index=1, content='Text')

        mocked_candidates = [
            # Candidate 1: MCQ with NO correct option
            {
                'question_type': 'mcq',
                'question_text': 'Which tree is balanced?',
                'marks': 1.0,
                'options': [
                    {'option_key': 'A', 'option_text': 'Opt 1', 'is_correct': False},
                    {'option_key': 'B', 'option_text': 'Opt 2', 'is_correct': False},
                ],
                'provenance': {'source_title': 'Notes', 'excerpt': 'Valid excerpt text passage for grounding'}
            },
            # Candidate 2: Negative marks exceed positive marks
            {
                'question_type': 'true_false',
                'question_text': 'An AVL tree is a BST?',
                'marks': 1.0,
                'negative_marks': 3.0,  # GREATER than marks
                'options': [
                    {'option_key': 'A', 'option_text': 'True', 'is_correct': True},
                    {'option_key': 'B', 'option_text': 'False', 'is_correct': False},
                ],
                'provenance': {'source_title': 'Notes', 'excerpt': 'Valid excerpt text passage for grounding'}
            }
        ]

        mock_instance = MockProvider.return_value
        mock_instance.generate_questions.return_value = (mocked_candidates, {})

        self.client.force_authenticate(user=self.teacher_alpha)
        job_url = reverse('question-generation-job-list')
        payload = {
            'institute': str(self.institute_alpha.id),
            'subject': str(self.subj_dsa.id),
            'source_document_ids': [str(doc.id)],
            'num_questions': 2
        }

        response = self.client.post(job_url, payload, format='json')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['status'], 'failed')
        self.assertEqual(Question.objects.count(), 0)

    @patch('apps.ai_generation.services.orchestrator.GoogleGeminiProvider')
    def test_provider_rate_limit_and_timeout_handling(self, MockProvider):
        doc = SourceDocument.objects.create(
            institute=self.institute_alpha,
            subject=self.subj_dsa,
            title='Notes',
            file_type='txt',
            processing_status='completed',
            extracted_text='Text'
        )
        SourceDocumentChunk.objects.create(document=doc, chunk_index=1, content='Text')

        mock_instance = MockProvider.return_value
        mock_instance.generate_questions.side_effect = AIProviderError("Gemini API rate limit or quota exceeded. Please retry later.")

        self.client.force_authenticate(user=self.teacher_alpha)
        job_url = reverse('question-generation-job-list')
        payload = {
            'institute': str(self.institute_alpha.id),
            'subject': str(self.subj_dsa.id),
            'source_document_ids': [str(doc.id)],
            'num_questions': 2
        }

        response = self.client.post(job_url, payload, format='json')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['status'], 'failed')
        self.assertIn("rate limit", response.data['error_message'])

    @patch('apps.ai_generation.services.orchestrator.GoogleGeminiProvider')
    def test_transaction_rollback_on_persistence_error(self, MockProvider):
        """
        If database persistence fails midway through candidate insertion,
        the transaction rollback prevents orphaned question records.
        """
        doc = SourceDocument.objects.create(
            institute=self.institute_alpha,
            subject=self.subj_dsa,
            title='Notes',
            file_type='txt',
            processing_status='completed',
            extracted_text='Text'
        )
        SourceDocumentChunk.objects.create(document=doc, chunk_index=1, content='Text')

        mocked_candidates = [
            {
                'question_type': 'mcq',
                'question_text': 'Valid candidate but persistence will simulate error',
                'marks': 1.0,
                'options': [
                    {'option_key': 'A', 'option_text': 'Opt A', 'is_correct': True},
                    {'option_key': 'B', 'option_text': 'Opt B', 'is_correct': False},
                ],
                'provenance': {'source_title': 'Notes', 'excerpt': 'Valid excerpt text passage'}
            }
        ]

        mock_instance = MockProvider.return_value
        mock_instance.generate_questions.return_value = (mocked_candidates, {})

        # Simulate exception during Question.objects.create
        with patch('apps.questions.models.Question.objects.create', side_effect=RuntimeError("Database write error")):
            self.client.force_authenticate(user=self.teacher_alpha)
            job_url = reverse('question-generation-job-list')
            payload = {
                'institute': str(self.institute_alpha.id),
                'subject': str(self.subj_dsa.id),
                'source_document_ids': [str(doc.id)],
                'num_questions': 1
            }
            response = self.client.post(job_url, payload, format='json')
            self.assertEqual(response.status_code, status.HTTP_201_CREATED)
            self.assertEqual(response.data['status'], 'failed')
            self.assertEqual(Question.objects.count(), 0)
