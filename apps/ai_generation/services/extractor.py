import io
import os
import pypdf
import docx


class DocumentExtractionError(Exception):
    """Custom exception raised when document extraction fails or format is invalid."""
    pass


SUPPORTED_EXTENSIONS = {
    'pdf': 'pdf',
    'docx': 'docx',
    'txt': 'txt',
}

MAX_FILE_SIZE_BYTES = 15 * 1024 * 1024  # 15 MB


def validate_document_file(file_obj, filename):
    """
    Validates file extension, size, and presence of content.
    Returns normalized extension ('pdf', 'docx', 'txt').
    """
    if not filename:
        raise DocumentExtractionError("Filename is missing.")

    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocumentExtractionError(
            f"Unsupported file format '{ext}'. Supported formats: {', '.join(SUPPORTED_EXTENSIONS.keys())}"
        )

    # Check size
    file_size = getattr(file_obj, 'size', None)
    if file_size is None:
        file_obj.seek(0, os.SEEK_END)
        file_size = file_obj.tell()
        file_obj.seek(0)

    if file_size <= 0:
        raise DocumentExtractionError("Uploaded file is empty (0 bytes).")
    if file_size > MAX_FILE_SIZE_BYTES:
        raise DocumentExtractionError(f"File size ({file_size} bytes) exceeds limit of 15MB.")

    return ext


def extract_text_from_pdf(file_bytes_io):
    """
    Extracts text page-by-page from PDF.
    Returns (full_text, total_pages, page_records).
    page_records is list of (page_num, text).
    """
    try:
        reader = pypdf.PdfReader(file_bytes_io)
        total_pages = len(reader.pages)
        if total_pages == 0:
            raise DocumentExtractionError("PDF file contains no pages.")

        page_records = []
        full_text_parts = []
        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ''
            page_text_clean = page_text.strip()
            if page_text_clean:
                page_records.append((idx + 1, page_text_clean))
                full_text_parts.append(f"--- Page {idx + 1} ---\n{page_text_clean}")

        full_text = "\n\n".join(full_text_parts).strip()
        if not full_text:
            raise DocumentExtractionError("Could not extract readable text from PDF (document may be empty or scanned images without OCR).")

        return full_text, total_pages, page_records

    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError(f"Malformed or corrupted PDF document: {str(exc)}") from exc


def extract_text_from_docx(file_bytes_io):
    """
    Extracts text paragraphs and table cells from DOCX.
    Returns (full_text, total_pages=1, page_records).
    """
    try:
        doc = docx.Document(file_bytes_io)
        text_parts = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                text_parts.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    text_parts.append(" | ".join(row_texts))

        full_text = "\n\n".join(text_parts).strip()
        if not full_text:
            raise DocumentExtractionError("DOCX document contains no readable text.")

        return full_text, 1, [(1, full_text)]

    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError(f"Malformed or corrupted DOCX document: {str(exc)}") from exc


def extract_text_from_txt(file_bytes_io):
    """
    Extracts text from plain text file with UTF-8 decoding and fallback.
    Returns (full_text, total_pages=1, page_records).
    """
    raw_data = file_bytes_io.read()
    try:
        text = raw_data.decode('utf-8')
    except UnicodeDecodeError:
        try:
            text = raw_data.decode('latin-1')
        except Exception as exc:
            raise DocumentExtractionError("Text file cannot be decoded (unsupported character encoding).") from exc

    full_text = text.strip()
    if not full_text:
        raise DocumentExtractionError("Plain text file is empty.")

    return full_text, 1, [(1, full_text)]


def extract_document_content(file_obj, filename):
    """
    Main extraction dispatcher.
    Returns:
        dict: {
            'file_type': str,
            'extracted_text': str,
            'total_pages': int,
            'page_records': list of (page_num, text)
        }
    """
    file_type = validate_document_file(file_obj, filename)
    file_obj.seek(0)
    data = file_obj.read()
    bytes_io = io.BytesIO(data)

    if file_type == 'pdf':
        full_text, total_pages, page_records = extract_text_from_pdf(bytes_io)
    elif file_type == 'docx':
        full_text, total_pages, page_records = extract_text_from_docx(bytes_io)
    elif file_type == 'txt':
        full_text, total_pages, page_records = extract_text_from_txt(bytes_io)
    else:
        raise DocumentExtractionError(f"Unsupported document format: {file_type}")

    return {
        'file_type': file_type,
        'extracted_text': full_text,
        'total_pages': total_pages,
        'page_records': page_records
    }


def chunk_extracted_content(page_records, max_chunk_chars=1800, overlap_chars=150):
    """
    Splits extracted text into clean, structured chunks while preserving page metadata.
    Returns list of dicts:
        [{'chunk_index': 1, 'page_number': 1, 'section_title': '...', 'content': '...', 'char_count': int}]
    """
    chunks = []
    chunk_index = 1

    for page_num, page_text in page_records:
        if len(page_text) <= max_chunk_chars:
            chunks.append({
                'chunk_index': chunk_index,
                'page_number': page_num,
                'section_title': f"Page {page_num}",
                'content': page_text,
                'char_count': len(page_text)
            })
            chunk_index += 1
        else:
            # Sub-chunk longer pages
            start = 0
            sub_part = 1
            while start < len(page_text):
                end = min(start + max_chunk_chars, len(page_text))
                # Break at paragraph or newline if near boundary
                if end < len(page_text):
                    last_newline = page_text.rfind('\n', start + max_chunk_chars - 300, end)
                    if last_newline > start + 300:
                        end = last_newline

                chunk_content = page_text[start:end].strip()
                if chunk_content:
                    chunks.append({
                        'chunk_index': chunk_index,
                        'page_number': page_num,
                        'section_title': f"Page {page_num} (Part {sub_part})",
                        'content': chunk_content,
                        'char_count': len(chunk_content)
                    })
                    chunk_index += 1
                    sub_part += 1

                start = end - overlap_chars if end < len(page_text) else end

    return chunks
